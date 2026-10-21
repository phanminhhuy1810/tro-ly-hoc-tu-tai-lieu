import copy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from pptx import Presentation

from study_backend.app import create_app
from study_backend.extract import extract
from study_backend.models import ChatReply, Lesson, LessonReading, Outline, QUESTION_DRAFTS
from study_backend.chapter_context import ChapterBrief

FIXTURE_SOURCE = b"IPv4 has 32 bits.\nA /26 block has 64 addresses.\nProcedure: (1) read the prefix; (2) calculate host bits."


class TestProvider:
    """Explicit test double. The running server never defaults to this provider."""
    calls = 0

    def __init__(self):
        self.instructions = []

    async def status(self):
        return {"provider": "test_double", "ready": True, "model": "test"}

    async def generate(self, schema, instruction, fragments):
        self.calls += 1
        self.instructions.append(instruction)
        refs = [f["id"] for f in fragments]
        if schema is ChapterBrief:
            return {'summary':'Phân tích chương từ nguồn test.', 'key_ideas':[fragments[-1]['text'][:240]],
                    'connections':[], 'misconceptions':[], 'question_targets':['Hiểu và áp dụng nguyên lý.'],
                    'source_refs':refs}
        if schema is ChatReply:
            return {"text": "Gợi ý từ nguồn test", "source_refs": refs, "needs_more_source": False}
        if schema is Outline:
            return {"title": "Khóa từ nguồn test", "units": [{"title": f"Bài {i}", "objective": "Đối chiếu nguồn", "source_refs": refs} for i in [1, 2]]}
        if schema in {Lesson, LessonReading, *QUESTION_DRAFTS.values()}:
            common = {"explanation": "Giải thích từ nguồn test", "source_refs": refs, "rubric": []}
            lesson = {
                "title": "Bài từ nguồn test",
                "sections": [{"title": "Đọc nguồn", "text": "Nội dung test có trích đoạn.", "source_refs": refs}],
                "questions": [
                    {**common, "kind": "mcq", "prompt": "Chọn A", "options": ["A", "B"], "left": [], "answer": "A"},
                    {**common, "kind": "cloze", "prompt": "Điền ____", "options": [], "left": [], "answer": "62"},
                    {**common, "kind": "matching", "prompt": "Ghép", "options": ["N", "B"], "left": ["Network", "Broadcast"], "answer": [0, 1]},
                    {**common, "kind": "ordering", "prompt": "Xếp", "options": ['"read the prefix"', '"calculate host bits"'], "left": [], "answer": [0, 1]},
                    {**common, "kind": "short_answer", "prompt": "Giải thích", "options": [], "left": [], "answer": "Đáp án tham khảo", "rubric": ["Tự đối chiếu"]},
                ],
                "cards": [{"front": "Network?", "back": "Địa chỉ đầu", "source_refs": refs}],
                "mindmap": [{"title": "Network", "text": "Địa chỉ đầu", "source_refs": refs}],
            }
            if '"language": "en"' in instruction:
                lesson['title']='Test lesson'
                lesson['sections']=[{'title':'Read source','text':'Test content with source citations.','source_refs':refs}]
                lesson['cards']=[{'front':'Network?','back':'The first address','source_refs':refs}]
                lesson['mindmap']=[{'title':'Network','text':'The first address','source_refs':refs}]
                for q in lesson['questions']:
                    q['prompt']={'mcq':'Choose A','cloze':'Fill ____','matching':'Match concepts','ordering':'Arrange steps','short_answer':'Explain the answer'}[q['kind']]
                    q['explanation']='Explanation from the test source.'
                    if q['kind']=='short_answer':q['rubric']=['Review the reasoning']
            if schema is LessonReading:
                return {k: v for k, v in lesson.items() if k != "questions"}
            if schema in QUESTION_DRAFTS.values():
                kind = next(k for k, v in QUESTION_DRAFTS.items() if v is schema)
                return {"question": next(q for q in lesson["questions"] if q["kind"] == kind), "reason": ""}
            return lesson
        raise RuntimeError("Unsupported test operation")


class SliceFixture(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        root = Path(self.temp.name)
        self.provider = TestProvider()
        # Domain tests isolate Coin; test_economy exercises enforcement explicitly.
        self.app = create_app(root / "data.sqlite", root / "sources", self.provider, extractor=lambda name, data: extract(name, data, ocr=False), economy_enabled=False)
        self.client = TestClient(self.app, base_url="http://127.0.0.1")
        self.client.post("/api/session").raise_for_status()

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def job(self, response):
        response.raise_for_status()
        job = self.client.get("/api/jobs/" + response.json()["job_id"]).json()
        self.assertEqual(job["status"], "succeeded", job)
        return job["result"]

    def source(self):
        data = FIXTURE_SOURCE
        source = self.client.post("/api/sources/upload", files={"file": ("notes.txt", data, "text/plain")}).json()
        self.assertFalse(source["reviewed"])
        self.assertEqual(self.client.get(f"/api/sources/{source['id']}/original").content, data)
        return source

    def preferences(self, source):
        return {"source_id": source["id"], "subject": "Mạng máy tính", "goal": "Học network", "fragment_ids": [f["id"] for f in source["fragments"]], "language": "vi"}

    def course(self):
        source = self.source()
        source = self.client.post(f"/api/sources/{source['id']}/review", json={"expected_revision": 1, "texts": [f["text"] for f in source["fragments"]]}).json()
        plan_id = self.job(self.client.post("/api/plans", json=self.preferences(source)))["plan_id"]
        approved = self.client.post(f"/api/plans/{plan_id}/approve", json={"expected_revision": 1})
        approved.raise_for_status()
        course_id = approved.json()["course_id"]
        self.job(self.client.post(f"/api/courses/{course_id}/generate", json={"expected_revision": 1, "unit_id": "unit-1"}))
        return self.client.get(f"/api/courses/{course_id}/editor").json()


class SliceTests(SliceFixture):
    def test_review_is_required_and_original_is_immutable(self):
        source = self.source()
        response = self.client.post("/api/plans", json=self.preferences(source))
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.provider.calls, 0)
        texts = [f["text"] for f in source["fragments"]]
        texts[0] += " User correction."
        body = {"expected_revision": 1, "texts": texts}
        edited = self.client.post(f"/api/sources/{source['id']}/review", json=body).json()
        self.assertEqual(edited["revision"], 2)
        self.assertTrue(edited["fragments"][0]["edited_by_user"])
        self.assertNotEqual(edited["fragments"][0]["id"], source["fragments"][0]["id"])
        self.assertEqual(edited["fragments"][0]["original_text"], source["fragments"][0]["text"])
        self.assertEqual(self.client.post(f"/api/sources/{source['id']}/review", json=body).status_code, 409)

    def test_other_session_cannot_read_source_course_jobs_or_apply(self):
        course = self.course()
        plan = self.client.get(f"/api/plans/{course['plan_id']}").json()
        other = TestClient(self.app, base_url="http://127.0.0.1")
        other.post("/api/session")
        self.assertEqual(other.get("/api/sources").json(), [])
        self.assertEqual(other.get("/api/courses").json(), [])
        for url in [f"/api/sources/{plan['source_id']}", f"/api/sources/{plan['source_id']}/original", f"/api/courses/{course['id']}", f"/api/plans/{course['plan_id']}"]:
            self.assertEqual(other.get(url).status_code, 404)
        self.assertEqual(other.post(f"/api/courses/{course['id']}/generate", json={"expected_revision": 2, "unit_id": "unit-1"}).status_code, 404)
        other.close()

    def test_attempt_pins_old_revision_and_zero_score_unlocks_next(self):
        course = self.course()
        base = f"/api/courses/{course['id']}"
        self.assertEqual(self.client.post(base + "/generate", json={"expected_revision": 2, "unit_id": "unit-2"}).status_code, 403)
        session = self.client.post(base + "/sessions", json={"expected_revision": 2, "unit_id": "unit-1"}).json()
        self.assertNotIn("answer", session["questions"][0])
        changed = copy.deepcopy(course["content"]["units"][0]["data"]["questions"][0])
        changed["answer"] = "B"
        proposal = self.client.post(base + "/proposals", json={"expected_revision": 2, "unit_id": "unit-1", "target": "question", "index": 0, "request": "Sửa đáp án", "replacement": changed}).json()
        preview = self.client.get("/api/proposals/" + proposal["proposal_id"]).json()
        self.assertEqual(preview["before"]["answer"], "A")
        self.assertEqual(self.client.get(base).json()["revision"], 2)
        applied = self.client.post(f"/api/proposals/{preview['id']}/apply", json={"expected_revision": 2})
        self.assertEqual(applied.json()["revision"], 3)
        values = ["B", "0", [1, 0], [1, 0], "Chưa hiểu"]
        for question, value in zip(session["questions"], values):
            response = self.client.post(f"/api/sessions/{session['id']}/answers", json={"item_id": question["id"], "answer": value})
            response.raise_for_status()
        result = self.client.post(f"/api/sessions/{session['id']}/complete").json()
        self.assertEqual((result["score"], result["fixed_total"], result["self_review_count"]), (0, 4, 1))
        self.assertEqual(result["course_revision"], 2)
        restored = self.client.get(f"/api/sessions/{session['id']}").json()
        self.assertEqual(restored["revision"], 2)
        self.assertEqual(len(restored["answers"]), 5)
        self.assertEqual(restored["result"], result)
        self.assertEqual(restored["questions"][0]["prompt"], session["questions"][0]["prompt"])
        self.assertEqual(self.client.post(f"/api/sessions/{session['id']}/complete").json(), result)
        self.assertTrue(self.client.get(base).json()["content"]["units"][1]["unlocked"])
        self.assertEqual(len(self.client.get(base + "/history").json()), 1)

    def test_invalid_matching_and_stale_proposals_are_rejected(self):
        course = self.course()
        base = f"/api/courses/{course['id']}"
        session = self.client.post(base + "/sessions", json={"expected_revision": 2, "unit_id": "unit-1"}).json()
        for value in [[0, 0], [True, False]]:
            self.assertEqual(self.client.post(f"/api/sessions/{session['id']}/answers", json={"item_id": session["questions"][2]["id"], "answer": value}).status_code, 422)
        changed = copy.deepcopy(course["content"]["units"][0]["data"]["cards"][0])
        changed["back"] = "Bản sửa"
        body = {"expected_revision": 2, "unit_id": "unit-1", "target": "card", "request": "Sửa thẻ", "replacement": changed}
        a = self.client.post(base + "/proposals", json=body).json()["proposal_id"]
        b = self.client.post(base + "/proposals", json=body).json()["proposal_id"]
        self.assertEqual(self.client.post(f"/api/proposals/{a}/apply", json={"expected_revision": 2}).status_code, 200)
        self.assertEqual(self.client.post(f"/api/proposals/{b}/apply", json={"expected_revision": 2}).status_code, 409)
        changed["source_refs"] = ["another-person-source"]
        body.update(expected_revision=3, replacement=changed)
        self.assertEqual(self.client.post(base + "/proposals", json=body).status_code, 422)

    def test_bad_inputs_and_cross_origin_are_not_accepted(self):
        self.assertEqual(self.client.post("/api/sources/text", json={"name": "x", "text": "test"}, headers={"Origin": "https://unrelated.example"}).status_code, 403)
        self.assertEqual(self.client.post("/api/sources/upload", files={"file": ("bad.docx", b"not a zip")}).status_code, 422)
        blank = PdfWriter(); blank.add_blank_page(width=100, height=100)
        data = BytesIO(); blank.write(data)
        response = self.client.post("/api/sources/upload", files={"file": ("scan.pdf", data.getvalue())})
        self.assertEqual(response.status_code, 422)
        self.assertIn("OCR", response.json()["detail"]["message"])
        failed_id = response.json()["detail"]["source_id"]
        self.assertEqual(self.client.get(f"/api/sources/{failed_id}/original").content, data.getvalue())
        self.assertEqual(self.client.get(f"/api/sources/{failed_id}").json()["status"], "failed")

    def test_oversized_request_is_rejected_before_upload_is_saved(self):
        response = self.client.post("/api/sources/upload", files={"file": ("large.txt", b"x" * (50 * 1024 * 1024 + 65537))})
        self.assertEqual(response.status_code, 413)
        self.assertEqual(list((Path(self.temp.name) / "sources").iterdir()), [])

    def test_failed_generation_keeps_source_and_approved_outline(self):
        source = self.source()
        reviewed = self.client.post(f"/api/sources/{source['id']}/review", json={"expected_revision": 1, "texts": [f["text"] for f in source["fragments"]]}).json()
        plan_id = self.job(self.client.post("/api/plans", json=self.preferences(reviewed)))["plan_id"]
        course_id = self.client.post(f"/api/plans/{plan_id}/approve", json={"expected_revision": 1}).json()["course_id"]
        previous_generate = self.provider.generate

        async def invalid_output(schema, instruction, fragments):
            result = await previous_generate(schema, instruction, fragments)
            if schema is QUESTION_DRAFTS["matching"]:
                result["question"]["source_refs"] = ["invented-fragment"]
            return result

        self.provider.generate = invalid_output
        response = self.client.post(f"/api/courses/{course_id}/generate", json={"expected_revision": 1, "unit_id": "unit-1"})
        failed = self.client.get("/api/jobs/" + response.json()["job_id"]).json()
        self.assertEqual(failed["status"], "failed")
        self.assertEqual(failed["attempts"], 1)
        self.assertEqual(failed['error']['code'],'SOURCE_REFERENCE_MISMATCH')
        self.assertIn('ngoài phạm vi nguồn',failed['error']['message'])
        self.assertIn("reading", failed["partial"])
        self.assertEqual(self.client.get(f"/api/plans/{plan_id}").json()["status"], "approved")
        course = self.client.get(f"/api/courses/{course_id}").json()
        self.assertEqual(course["revision"], 1)
        self.assertIsNone(course["content"]["units"][0]["data"])
        self.assertEqual(self.client.get(f"/api/sources/{source['id']}/original").content, FIXTURE_SOURCE)
        other = TestClient(self.app, base_url="http://127.0.0.1"); other.post("/api/session")
        self.assertEqual(other.get("/api/jobs/" + failed["id"]).status_code, 404)
        other.close()
        self.provider.generate = previous_generate
        prior_calls = self.provider.calls
        self.job(self.client.post(f"/api/courses/{course_id}/generate", json={"expected_revision": 1, "unit_id": "unit-1", "resume_job_id": failed["id"]}))
        self.assertEqual(self.provider.calls - prior_calls, 3, "Retry must reuse reading and the two successful question types")

    def test_proposal_is_private_and_discard_does_not_change_content(self):
        course = self.course()
        changed = copy.deepcopy(course["content"]["units"][0]["data"]["cards"][0])
        changed["back"] = "Bản đề xuất chưa áp dụng"
        proposal_id = self.client.post(f"/api/courses/{course['id']}/proposals", json={"expected_revision": 2, "unit_id": "unit-1", "target": "card", "request": "Sửa thử", "replacement": changed}).json()["proposal_id"]
        other = TestClient(self.app, base_url="http://127.0.0.1"); other.post("/api/session")
        self.assertEqual(other.get("/api/proposals/" + proposal_id).status_code, 404)
        self.assertEqual(other.post(f"/api/proposals/{proposal_id}/apply", json={"expected_revision": 2}).status_code, 404)
        other.close()
        self.client.post(f"/api/proposals/{proposal_id}/discard").raise_for_status()
        self.assertEqual(self.client.post(f"/api/proposals/{proposal_id}/apply", json={"expected_revision": 2}).status_code, 409)
        self.assertEqual(self.client.get(f"/api/courses/{course['id']}/editor").json()["content"], course["content"])

    def test_docx_and_pptx_have_real_locators(self):
        document = Document(); document.add_heading("Môn thử", 1); document.add_paragraph("Đoạn cần đọc")
        table = document.add_table(rows=1, cols=1); table.cell(0, 0).text = "Bảng ở giữa"
        document.add_paragraph("Đoạn sau bảng")
        data = BytesIO(); document.save(data)
        fragments, _ = extract("notes.docx", data.getvalue())
        self.assertEqual(fragments[1]["locator"]["paragraph"], 2)
        self.assertEqual(fragments[1]["locator"]["heading"], "Môn thử")
        self.assertEqual([f["text"] for f in fragments], ["Môn thử", "Đoạn cần đọc", "Bảng ở giữa", "Đoạn sau bảng"])
        self.assertEqual(fragments[2]["locator"]["table"], 1)
        self.assertEqual(fragments[3]["locator"]["paragraph"], 3)
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text = "Slide thử"
        data = BytesIO(); deck.save(data)
        fragments, _ = extract("notes.pptx", data.getvalue())
        self.assertEqual(fragments[0]["locator"]["slide"], 1)

    def test_pdf_text_is_located_and_blank_pages_are_reported(self):
        writer = PdfWriter(); page = writer.add_blank_page(width=300, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject(); stream.set_data(b"BT /F1 12 Tf 30 150 Td (IPv4 has 32 bits.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.add_blank_page(width=300, height=200)
        data = BytesIO(); writer.write(data)
        fragments, warnings = extract("text.pdf", data.getvalue(), ocr=False)
        self.assertEqual(fragments[0]["text"], "IPv4 has 32 bits.")
        self.assertEqual(fragments[0]["locator"]["page"], 1)
        self.assertTrue(any("Trang 2" in warning and "OCR" in warning for warning in warnings))

    def test_chat_keeps_context_across_revision_and_remains_private(self):
        course = self.course()
        base = f"/api/courses/{course['id']}"
        self.job(self.client.post(base + "/chat", json={"unit_id": "unit-1", "message": "Hỏi về /26 ở bản cũ"}))
        changed = copy.deepcopy(course["content"]["units"][0]["data"]["cards"][0]); changed["back"] = "Bản chỉnh tay"
        proposal = self.client.post(base + "/proposals", json={"expected_revision": 2, "unit_id": "unit-1", "target": "card", "request": "Sửa thẻ", "replacement": changed}).json()
        self.client.post(f"/api/proposals/{proposal['proposal_id']}/apply", json={"expected_revision": 2}).raise_for_status()
        self.job(self.client.post(base + "/chat", json={"unit_id": "unit-1", "message": "Cho gợi ý bước tiếp"}))
        self.assertIn("Hỏi về /26 ở bản cũ", self.provider.instructions[-1])
        self.assertIn('"course_revision": 3', self.provider.instructions[-1])
        self.assertEqual([m["revision"] for m in self.client.get(base + "/chat").json()], [2, 2, 3, 3])
        other = TestClient(self.app, base_url="http://127.0.0.1"); other.post("/api/session")
        self.assertEqual(other.get(base + "/chat").status_code, 404); other.close()


if __name__ == "__main__":
    unittest.main()
