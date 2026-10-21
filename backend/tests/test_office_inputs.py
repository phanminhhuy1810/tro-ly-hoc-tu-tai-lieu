"""Generated OOXML fixtures, real package/pixel processing, explicit OCR/AI doubles."""
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches
from fastapi.testclient import TestClient
from lxml import etree
from PIL import Image
from pptx import Presentation

from study_backend.app import create_app
from study_backend.extract import extract
from study_backend.office_inputs import read_office,raster_asset,read_with_ocr
from study_backend.office_math import M,transcribe,UNREADABLE
from test_slice import TestProvider


def picture():
    image=Image.new('RGB',(200,100),'red');image.paste('blue',(100,0,200,100))
    output=BytesIO();image.save(output,format='PNG');return output.getvalue()


def math_xml(contents):return etree.fromstring(('<m:oMath xmlns:m="'+M+'">'+contents+'</m:oMath>').encode())
def run(text):return '<m:r><m:t>'+escape(text)+'</m:t></m:r>'


def word_fixture(images=True,maths=None,external=False,image_only=False):
    doc=Document()
    if not image_only:doc.add_paragraph('Native paragraph before the image.')
    if maths is not None:
        p=doc.add_paragraph('Formula: ');p._p.append(math_xml(maths));p.add_run(' after formula.')
    if images:
        inline=doc.add_picture(BytesIO(picture()),width=Inches(2))
        fill=inline._inline.graphic.graphicData.pic.blipFill
        crop=OxmlElement('a:srcRect');crop.set('l','50000');fill.insert(1,crop)
        if external:
            blip=fill.find('{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
            blip.attrib.clear();blip.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link',doc.part.relate_to('https://external.invalid/do-not-fetch',RT.IMAGE,is_external=True))
        if not image_only:
            table=doc.add_table(rows=1,cols=1);table.cell(0,0).paragraphs[0].add_run().add_picture(BytesIO(picture()),width=Inches(1))
    if not image_only:doc.add_paragraph('Native paragraph after the image.')
    out=BytesIO();doc.save(out);return out.getvalue()


class OfficeParsing(unittest.TestCase):
    def test_word_occurrences_have_stable_distinct_locations_and_only_visible_crop(self):
        raw=word_fixture();blocks,assets,_=read_office('fixture.docx',raw)
        self.assertEqual(len(assets),2);self.assertNotEqual(assets[0].id,assets[1].id)
        self.assertEqual(assets[0].public()['sha256'],assets[1].public()['sha256'])
        self.assertEqual(assets[0].locator['paragraph'],2);self.assertEqual(assets[1].locator['table'],1)
        self.assertEqual(assets[1].locator['cell'],1)
        self.assertEqual([a.id for a in read_office('fixture.docx',raw)[1]],[a.id for a in assets])
        image,dimensions=raster_asset(assets[0]);decoded=Image.open(BytesIO(image))
        self.assertEqual(decoded.size,(100,100));self.assertEqual(decoded.getpixel((0,0)),(0,0,255))
        self.assertEqual(decoded.getextrema(),((0,0),(0,0),(255,255)))
        self.assertEqual(dimensions['scope'],'embedded_image_crop')
        self.assertEqual([b['text'].strip() for b in blocks if b['text'].strip()],['Native paragraph before the image.','Native paragraph after the image.'])

    def test_grouped_powerpoint_picture_is_located_without_claiming_a_slide_render(self):
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);group=slide.shapes.add_group_shape()
        p=group.shapes.add_picture(BytesIO(picture()),0,0);p.crop_left=.5
        out=BytesIO();deck.save(out)
        _,assets,_=read_office('fixture.pptx',out.getvalue())
        self.assertEqual(assets[0].locator['slide'],1);self.assertEqual(assets[0].locator['shape_path'],[1,1])
        image,dimensions=raster_asset(assets[0]);self.assertEqual(Image.open(BytesIO(image)).getpixel((0,0)),(0,0,255))
        self.assertEqual(dimensions['scope'],'embedded_image_crop')

    def test_external_link_is_not_fetched_or_exposed_as_an_embedded_image(self):
        with patch('urllib.request.urlopen',side_effect=AssertionError('must not fetch')):
            _,assets,warnings=read_office('fixture.docx',word_fixture(external=True))
            public=assets[0].public()
            self.assertFalse(public['embedded']);self.assertFalse(public['raster_review_supported'])
            self.assertNotIn('https://external.invalid',str(public))
            with self.assertRaises(ValueError):raster_asset(assets[0])
        self.assertTrue(any('liên kết ngoài' in w for w in warnings))

    def test_omml_fraction_superscript_and_integral_preserve_structure_and_inline_order(self):
        fraction='<m:f><m:num>'+run('1')+'</m:num><m:den>'+run('3')+'</m:den></m:f>'
        integral='<m:nary><m:naryPr><m:chr m:val="∫"/></m:naryPr><m:sub>'+run('0')+'</m:sub><m:sup>'+run('1')+'</m:sup><m:e><m:sSup><m:e>'+run('x')+'</m:e><m:sup>'+run('2')+'</m:sup></m:sSup>'+run('dx')+'</m:e></m:nary>'+run('=')+fraction
        fragments,_=extract('math.docx',word_fixture(images=False,maths=integral),ocr=False)
        formula=next(f for f in fragments if f['locator'].get('office_math'))
        self.assertIn(r'\int_{0}^{1} {x}^{2}dx=\frac{1}{3}',formula['text'])
        self.assertTrue(formula['text'].startswith('Formula: '));self.assertTrue(formula['text'].endswith(' after formula.'))
        self.assertFalse(formula['locator'].get('unreadable_math',False))
        self.assertFalse(formula['locator']['office_math'][0]['accuracy_verified'])

    def test_unsupported_omml_is_marked_instead_of_flattened_into_a_false_formula(self):
        node=math_xml('<m:unknown>'+run('x')+'</m:unknown>')
        text,info=transcribe(node,1)
        self.assertEqual(text,UNREADABLE);self.assertFalse(info['supported'])
        self.assertIn('element_unknown',info['reason'])
        fragments,_=extract('math.docx',word_fixture(images=False,maths='<m:unknown>'+run('x')+'</m:unknown>'),ocr=False)
        self.assertTrue(next(f for f in fragments if UNREADABLE in f['text'])['locator']['unreadable_math'])

    def test_native_image_ocr_has_provenance_and_does_not_replace_native_paragraphs(self):
        calls=[]
        def recognize(name,image,**options):
            calls.append(options)
            return [{'text':'Image text from explicit OCR double','locator':{'bounds':[.1,.2,.3,.4]},'recognizer_score':.9}],['OCR requires review']
        blocks,_=read_with_ocr('fixture.docx',word_fixture(),True,recognize)
        recovered=[f for f in blocks if f.get('provenance')]
        self.assertEqual(len(recovered),2);self.assertEqual(len(calls),2)
        self.assertTrue(all(0<call['timeout_seconds']<=40 for call in calls))
        self.assertTrue(all(f['needs_review'] and f['locator']['asset_id'] for f in recovered))
        self.assertEqual(len([f for f in blocks if f['text'].startswith('Native paragraph')]),2)

    def test_invalid_crop_is_reported_and_never_sends_the_uncropped_picture(self):
        _,assets,_=read_office('fixture.docx',word_fixture())
        assets[0].crop=(.8,0,.3,0)
        with self.assertRaises(ValueError):raster_asset(assets[0])


class OfficeAPI(unittest.TestCase):
    def setUp(self):
        self.tmp=TemporaryDirectory();root=Path(self.tmp.name)
        class Vision(TestProvider):
            async def transcribe_image(provider,image):
                provider.calls+=1
                self.assertEqual(Image.open(BytesIO(image)).getpixel((0,0)),(0,0,255))
                return {'text':'AI transcription for chosen image','uncertain_spans':[],'model':'explicit_test_double'}
        self.provider=Vision();self.app=create_app(root/'data.sqlite',root/'sources',self.provider,extractor=lambda name,data:extract(name,data,ocr=False),economy_enabled=False)
        self.client=TestClient(self.app,base_url='http://127.0.0.1');self.client.post('/api/session')
        self.raw=word_fixture();self.source=self.client.post('/api/sources/upload',files={'file':('fixture.docx',self.raw)}).json();self.path='/api/sources/'+self.source['id']
        self.assets=self.client.get(self.path+'/assets').json()['assets']

    def tearDown(self):self.client.close();self.tmp.cleanup()

    def proposal(self):
        response=self.client.post(self.path+'/vision-review',json={'expected_revision':1,'asset_id':self.assets[0]['id']},headers={'Idempotency-Key':'chosen-office-picture'})
        response.raise_for_status();job=self.client.get('/api/jobs/'+response.json()['job_id']).json()
        self.assertEqual(job['status'],'succeeded',job)
        return self.client.get('/api/source-proposals/'+job['result']['source_proposal_id']).json()

    def test_image_inventory_renderer_and_vision_preview_are_owner_scoped(self):
        image_path=self.path+'/assets/'+self.assets[0]['id']+'/image'
        self.assertEqual(self.client.get(image_path).status_code,200)
        proposal=self.proposal();self.assertEqual(proposal['asset_id'],self.assets[0]['id'])
        self.assertEqual(self.client.get(self.path).json()['revision'],1)
        foreign=TestClient(self.app,base_url='http://127.0.0.1');foreign.post('/api/session')
        for path in [self.path+'/assets',image_path,'/api/source-proposals/'+proposal['id']]:self.assertEqual(foreign.get(path).status_code,404)
        self.assertEqual(foreign.post(self.path+'/vision-review',json={'expected_revision':1,'asset_id':self.assets[0]['id']}).status_code,404);foreign.close()
        self.assertEqual(self.provider.calls,1)

    def test_manual_edit_preserves_ai_proposal_and_only_applies_to_chosen_occurrence(self):
        proposal=self.proposal();parent='/api/source-proposals/'+proposal['id']
        edited=self.client.post(parent+'/edit-preview',json={'expected_revision':1,'text':'Manual corrected formula: x^2.'});edited.raise_for_status();edited=edited.json()
        self.assertEqual(self.provider.calls,1);self.assertEqual(self.client.get(parent).json()['text'],proposal['text'])
        self.assertEqual(self.client.get(self.path).json()['revision'],1)
        applied=self.client.post('/api/source-proposals/'+edited['id']+'/apply',json={'expected_revision':1});applied.raise_for_status();source=applied.json()
        self.assertFalse(source['reviewed']);self.assertEqual(source['revision'],2)
        image=next(f for f in source['fragments'] if f['locator'].get('asset_id')==self.assets[0]['id'])
        self.assertEqual(image['text'],'Manual corrected formula: x^2.');self.assertEqual(image['provenance']['kind'],'manual_transcription')
        self.assertEqual(image['provenance']['parent_proposal_id'],proposal['id'])
        self.assertEqual([f['text'] for f in source['fragments'] if not f['locator'].get('asset_id')],[f['text'] for f in self.source['fragments']])
        self.assertEqual(self.client.post(parent+'/apply',json={'expected_revision':2}).status_code,409)
        self.assertEqual(self.client.get(self.path+'/original').content,self.raw)
        self.assertEqual(self.client.get(self.path+'?revision=1').json()['fragments'],self.source['fragments'])

    def test_image_only_failed_extraction_can_be_recovered_by_manual_preview_without_ai(self):
        raw=word_fixture(image_only=True);response=self.client.post('/api/sources/upload',files={'file':('image-only.docx',raw)})
        self.assertEqual(response.status_code,422);path='/api/sources/'+response.json()['detail']['source_id']
        asset=self.client.get(path+'/assets').json()['assets'][0]
        response=self.client.post(path+'/assets/'+asset['id']+'/transcription-preview',json={'expected_revision':0,'text':'Manual image-only content.'});response.raise_for_status();proposal=response.json()
        self.assertEqual(self.client.get(path).json()['revision'],0)
        source=self.client.post('/api/source-proposals/'+proposal['id']+'/apply',json={'expected_revision':0});source.raise_for_status();source=source.json()
        self.assertEqual(source['revision'],1);self.assertFalse(source['reviewed']);self.assertEqual(self.provider.calls,0)
        self.assertEqual(self.client.get(path+'/original').content,raw)

    def test_unknown_math_cannot_reach_ai_until_explicitly_corrected_or_excluded(self):
        raw=word_fixture(images=False,maths='<m:unknown>'+run('x')+'</m:unknown>')
        source=self.client.post('/api/sources/upload',files={'file':('unknown.docx',raw)}).json();path='/api/sources/'+source['id']
        source=self.client.post(path+'/review',json={'expected_revision':1,'texts':[f['text'] for f in source['fragments']]}).json()
        bad=next(f for f in source['fragments'] if f['locator'].get('unreadable_math'))
        prefs={'source_id':source['id'],'subject':'Math','goal':'Read formula','fragment_ids':[bad['id']],'language':'en'}
        self.assertEqual(self.client.post('/api/plans',json=prefs).status_code,422);self.assertEqual(self.provider.calls,0)
        texts=[f['text'].replace(UNREADABLE,'f = x + y') for f in source['fragments']]
        corrected=self.client.post(path+'/review',json={'expected_revision':2,'texts':texts});corrected.raise_for_status();source=corrected.json()
        fixed=next(f for f in source['fragments'] if f['locator'].get('manual_math_correction'));prefs['fragment_ids']=[fixed['id']]
        response=self.client.post('/api/plans',json=prefs);response.raise_for_status()
        self.assertEqual(self.provider.calls,1)

    def test_external_asset_rejects_ai_but_accepts_manual_scope_without_fetching(self):
        source=self.client.post('/api/sources/upload',files={'file':('linked.docx',word_fixture(external=True))}).json();path='/api/sources/'+source['id']
        asset=self.client.get(path+'/assets').json()['assets'][0]
        self.assertEqual(self.client.post(path+'/vision-review',json={'expected_revision':1,'asset_id':asset['id']}).status_code,422)
        response=self.client.post(path+'/assets/'+asset['id']+'/transcription-preview',json={'expected_revision':1,'text':'User-provided visible image text.'})
        self.assertEqual(response.status_code,200);self.assertEqual(self.provider.calls,0)

    def test_failed_office_package_lists_a_recoverable_error_not_a_server_crash(self):
        response=self.client.post('/api/sources/upload',files={'file':('broken.docx',b'not an Office package')})
        self.assertEqual(response.status_code,422)
        path='/api/sources/'+response.json()['detail']['source_id']
        self.assertEqual(self.client.get(path+'/assets').status_code,422)
        self.assertEqual(self.client.get(path+'/original').content,b'not an Office package')
