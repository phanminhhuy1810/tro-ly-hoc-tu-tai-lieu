"""Self-authored Office development fixtures, not instructor or held-out materials."""
from io import BytesIO
import json
from pathlib import Path
import sys
import time
from zipfile import ZipFile, ZIP_STORED

from docx import Document
from docx.shared import Inches
from lxml import etree
from pptx import Presentation
from pptx.util import Inches as SlideInches

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'backend'))
from study_backend.extract import extract
from study_backend.office_inputs import read_office
from study_backend.office_math import M
HERE=Path(__file__).resolve().parent

def run(text):return '<m:r><m:t>'+text+'</m:t></m:r>'
def fixtures():
    doc=Document();doc.add_heading('Office parser development fixture',0)
    doc.add_paragraph('Self-authored QA. Not a lecture, student work or held-out benchmark.')
    for name,title in [('networking','Networking'),('calculus','Calculus'),('psychology','Psychology')]:
        doc.add_heading(title,1);doc.add_paragraph('Native text before the embedded image: '+title)
        doc.add_picture(str(ROOT/'experiments/03-ocr-inputs/fixtures'/f'{name}.png'),width=Inches(6))
        doc.add_paragraph('Native text after the embedded image: '+title)
    doc.add_heading('Native Office Math',1)
    p=doc.add_paragraph('Exact native formula: ')
    equation='<m:nary><m:naryPr><m:chr m:val="∫"/></m:naryPr><m:sub>'+run('0')+'</m:sub><m:sup>'+run('1')+'</m:sup><m:e><m:sSup><m:e>'+run('x')+'</m:e><m:sup>'+run('2')+'</m:sup></m:sSup>'+run('dx')+'</m:e></m:nary>'+run('=')+'<m:f><m:num>'+run('1')+'</m:num><m:den>'+run('3')+'</m:den></m:f>'
    p._p.append(etree.fromstring(('<m:oMath xmlns:m="'+M+'">'+equation+'</m:oMath>').encode()))
    p.add_run(' end of native formula.')
    doc.save(HERE/'fixtures/office-dev.docx')
    deck=Presentation();deck.slide_width=SlideInches(12);deck.slide_height=SlideInches(7)
    for name,title in [('networking','Networking'),('calculus','Calculus'),('psychology','Psychology')]:
        slide=deck.slides.add_slide(deck.slide_layouts[6]);box=slide.shapes.add_textbox(SlideInches(.5),SlideInches(.2),SlideInches(11),SlideInches(.8));box.text_frame.text=title+' - self-authored parser QA'
        group=slide.shapes.add_group_shape()
        group.shapes.add_picture(str(ROOT/'experiments/03-ocr-inputs/fixtures'/f'{name}.png'),SlideInches(.5),SlideInches(1.3),width=SlideInches(11))
    deck.save(HERE/'fixtures/office-dev.pptx')
    network=Document();network.add_heading('Networking upload-capacity fixture',0)
    network.add_paragraph('Self-authored QA, not a real lecture. IPv4 has 32 bits. A /26 block has 64 addresses and 62 usable host addresses.')
    network.add_picture(str(ROOT/'experiments/03-ocr-inputs/fixtures/networking.png'),width=Inches(6))
    out=BytesIO();network.save(out)
    capacity_dir=ROOT/'backend/runtime/office-capacity-fixtures';capacity_dir.mkdir(parents=True,exist_ok=True)
    with ZipFile(BytesIO(out.getvalue())) as original, ZipFile(capacity_dir/'network-large.docx','w',ZIP_STORED) as target:
        for entry in original.infolist():target.writestr(entry.filename,original.read(entry))
        target.writestr('developer-capacity-padding.bin',b'0'*(13*1024*1024))
    facts=['IPv4 has 32 bits.','Prefix /26 leaves 6 host bits.','A /26 block contains 64 addresses.','192.168.1.70/26 belongs to network 192.168.1.64.','The broadcast is 192.168.1.127.','Usable hosts are 192.168.1.65 through 192.168.1.126, or 62 addresses.']
    (HERE/'fixtures/network-notes.txt').write_text('\n'.join(f'Development excerpt {i+1}: '+facts[i%len(facts)] for i in range(80)))

def main():
    fixtures();report={'purpose':'Self-authored development parser probe, not a model-quality benchmark','cases':[]}
    for filename in ['office-dev.docx','office-dev.pptx']:
        raw=(HERE/'fixtures'/filename).read_bytes();start=time.monotonic()
        blocks,assets,notes=read_office(filename,raw)
        fragments,warnings=extract(filename,raw)
        report['cases'].append({'filename':filename,'elapsed_seconds':round(time.monotonic()-start,3),'baseline_native_paragraph_text':('\n'.join(p.text for p in Document(BytesIO(raw)).paragraphs) if filename.endswith('.docx') else '\n'.join(shape.text for slide in Presentation(BytesIO(raw)).slides for shape in slide.shapes if shape.has_text_frame)), 'native_blocks':blocks,'image_occurrences':[a.public() for a in assets],'fragments_with_ocr':fragments,'warnings':warnings})
    (HERE/'results/native-office-dev.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print([(c['filename'],len(c['image_occurrences']),len(c['fragments_with_ocr']),c['elapsed_seconds']) for c in report['cases']])
if __name__=='__main__':main()
