"""Self-authored development fixtures; outputs stay in private runtime, not the handoff."""
import json
from pathlib import Path

from PIL import Image
from pypdf import PdfReader, PdfWriter, PageObject
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
from pptx import Presentation
from pptx.util import Inches

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'backend/runtime/ocr-continuation-fixtures'


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    source=PdfReader(ROOT/'experiments/03-ocr-inputs/fixtures/networking-scan.pdf').pages[0]
    width,height=float(source.mediabox.width),float(source.mediabox.height)
    writer=PdfWriter()
    for number in range(1,48):
        page=PageObject.create_blank_page(width=width,height=height+40)
        page.merge_page(source)
        overlay=PageObject.create_blank_page(width=width,height=height+40)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),
                               NameObject('/BaseFont'):NameObject('/Helvetica')})
        overlay[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        stream=DecodedStreamObject();stream.set_data(f'BT /F1 12 Tf 24 {height+14} Td (QA native page {number}) Tj ET'.encode())
        overlay[NameObject('/Contents')]=stream;page.merge_page(overlay);writer.add_page(page)
    with (OUT/'networking-47-pages.pdf').open('wb') as file:writer.write(file)

    image=Image.open(ROOT/'experiments/03-ocr-inputs/fixtures/networking.png').convert('RGB')
    frames=[image.copy() for _ in range(46)]+[Image.new('RGB',image.size,'white')]
    frames[0].save(OUT/'networking-47-frames.tiff',format='TIFF',compression='tiff_lzw',
                   save_all=True,append_images=frames[1:])
    presentation=Presentation();presentation.slide_width=Inches(10);presentation.slide_height=Inches(6)
    for number in range(1,24):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(.3),Inches(.1),Inches(9),Inches(.4)).text=f'QA native slide {number}'
        slide.shapes.add_picture(str(ROOT/'experiments/03-ocr-inputs/fixtures/networking.png'),
                                 Inches(.3),Inches(.8),width=Inches(9))
    presentation.save(OUT/'networking-23-slides.pptx')
    print(json.dumps({'fixture_kind':'self_authored_development_not_held_out',
                      'files':{path.name:path.stat().st_size for path in OUT.iterdir()}},indent=2))


if __name__=='__main__':main()
